"""Exports are built only from an explicit user-confirmed result."""
import io
import json
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from rfq_inbox.review import columns, derive_issues


def build_exports(case, confirmed):
    if confirmed.get("origin") != "user_confirmed_review":
        raise ValueError("Confirm and save the review before exporting")
    missing, conflicts = derive_issues(confirmed)
    result = dict(confirmed, missing_information=missing, conflicts=conflicts)
    payload = {"case_id": case["case_id"], "result": result,
               "sources": [{"document_id": d["document_id"], "name": d["display_name"],
                            "source_files": d["source_files"]} for d in case["documents"]]}
    book = Workbook()
    sheet = book.active
    sheet.title = "Reviewed items"
    headers = columns(confirmed["document_type"])
    sheet.append(headers)
    for item in confirmed["items"]:
        sheet.append([item.get(key, {}).get("value") for key in headers])
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # Text identifiers and formula-looking untrusted strings stay literal.
                cell.number_format = "@"
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="183A55")
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 23
    binary = io.BytesIO()
    book.save(binary)
    title = "Invoice Review Summary" if confirmed["document_type"] == "invoice" else "RFQ Summary"
    summary = [title, "Case: " + case["case_id"], "Based on user-confirmed review values.", ""]
    if any(d.get("metadata", {}).get("historical") for d in case["documents"]):
        summary.append("HISTORICAL DEMONSTRATION MATERIAL - not a current procurement request.\n")
    if confirmed["document_type"] == "rfq_template":
        summary.append("Blank RFQ template structure only. No customer order was created.")
    for key, f in confirmed["fields"].items():
        summary.append(key.replace("_", " ").title() + ": " + (str(f["value"]) if f["value"] is not None else "Not provided"))
    summary += ["", "Items:"]
    for i, item in enumerate(confirmed["items"], 1):
        summary.append(str(i) + ". " + "; ".join(key + "=" + str(item[key]["value"]) for key in headers))
    summary += ["", "Missing information:"] + ["- " + x.replace("_", " ") for x in missing]
    summary += ["", "Unresolved conflicts:"] + ["- " + x for x in conflicts]
    summary += ["", "Review notes:"] + confirmed.get("review_notes", [])
    draft = ["Subject: Request for Missing Information", "", "Hello,", ""]
    if missing or conflicts:
        draft += ["To complete our review, could you please confirm the following?", ""]
        draft += ["- Please provide " + x.replace("_", " ") + "." for x in missing]
        draft += ["- Please resolve the conflicting information: " + x + "." for x in conflicts]
    else:
        draft += ["There are no currently missing or conflicting fields in the confirmed review. No clarification request is needed."]
    draft += ["", "Thank you.", ""]
    return {"review.json": json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8"),
            "items.xlsx": binary.getvalue(), "rfq-summary.txt": "\n".join(summary).encode("utf-8"),
            "request-for-missing-information.txt": "\n".join(draft).encode("utf-8")}


def save_review(case, confirmed, original_prediction=None):
    root = Path(case["storage_dir"])
    audit = {"confirmed_at_utc": datetime.now(timezone.utc).isoformat(), "result": confirmed,
             "original_prediction_file": original_prediction}
    (root / "review.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    target = root / "exports"
    target.mkdir(exist_ok=True)
    for name, content in build_exports(case, confirmed).items():
        (target / name).write_bytes(content)
    return audit
