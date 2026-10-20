"""Functional regressions; any model responses here are explicit test doubles."""
import copy
import io
import json
import os
import sys
import tempfile
import threading
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from unittest.mock import patch

PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT))
from rfq_inbox.exports import build_exports, save_review
from rfq_inbox.model import OpenAICompatibleProvider, ModelNotConfigured
from rfq_inbox.parsing import parse_case, source_units
from rfq_inbox.review import (blank_result, columns, confirm_review, derive_issues,
                              evidence_check, field, verify_prediction)


class InboxChecks(unittest.TestCase):
    def setUp(self):
        self.base = (PROJECT / "data/local-inbox/test-runs").resolve()
        self.base.mkdir(parents=True, exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(prefix="test-", dir=str(self.base))
        self.root = Path(self.temp.name).resolve()
        self.root.relative_to(self.base)

    def tearDown(self):
        self.root.resolve().relative_to(self.base)
        self.temp.cleanup()

    def example(self, slug, extra=False):
        path = PROJECT / "data/synthetic/rfq-inbox" / slug / "inputs/request.eml"
        files = [("request.eml", path.read_bytes())]
        if extra:
            files += [("items-renamed.xlsx", (path.parent / "items.xlsx").read_bytes())]
        return parse_case(files, storage=self.root)

    def test_case_relationship_duplicate_attachment_and_conflicting_sources(self):
        case = self.example("c-quantity-conflict", extra=True)
        self.assertEqual(len(case["documents"]), 2)
        self.assertEqual(len(case["emails"]), 1)
        self.assertEqual(len(case["emails"][0]["attachment_ids"]), 1)
        self.assertTrue(any(a["duplicate"] for a in case["attachments"]))
        self.assertEqual(len(case["source_conflicts"]), 1)
        conflict = case["source_conflicts"][0]
        self.assertEqual({c["value"] for c in conflict["candidates"]}, {"100", "120"})
        for candidate in conflict["candidates"]:
            for e in candidate["evidence"]:
                self.assertTrue(evidence_check(e, case)[0])
        wrong = dict(conflict["candidates"][1]["evidence"][-1], quote="999999")
        self.assertFalse(evidence_check(wrong, case)[0])

    def test_missing_material_is_absent_and_runtime_does_not_open_answers(self):
        original = Path.open
        def guarded(path, *args, **kwargs):
            if path.name in ("expected.json", "gold.jsonl") or "references" in path.parts:
                raise AssertionError("Runtime attempted to read a reference answer")
            return original(path, *args, **kwargs)
        with patch.object(Path, "open", guarded):
            case = self.example("b-missing-material")
            texts = "\n".join(u["text"] for u in source_units(case))
            self.assertNotIn("316 stainless steel", texts)
            self.assertEqual(case["source_conflicts"], [])
            self.assertTrue(all(d["processing_status"] == "ok" for d in case["documents"]))
            self.assertTrue(all(f["value"] is None for f in blank_result("rfq")["fields"].values()))

    def test_one_email_keeps_multiple_distinct_attachment_documents(self):
        from email.message import EmailMessage
        from openpyxl import load_workbook
        path = PROJECT / "data/synthetic/rfq-inbox/a-complete/inputs/items.xlsx"
        book = load_workbook(io.BytesIO(path.read_bytes()))
        book.active["A4"] = "0008-A"
        content = io.BytesIO(); book.save(content)
        message = EmailMessage()
        message["From"] = "buyer@example.com"; message["To"] = "quotes@example.com"
        message["Subject"] = "SYNTHETIC DEMO - two attachments"
        message.set_content("Please review both attached schedules.")
        for name, data in (("one.xlsx", path.read_bytes()), ("two.xlsx", content.getvalue())):
            message.add_attachment(data, maintype="application", subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=name)
        case = parse_case([("multiple.eml", message.as_bytes())], storage=self.root)
        self.assertEqual(len(case["documents"]), 3)
        self.assertEqual(len(case["emails"][0]["attachment_ids"]), 2)
        self.assertEqual(len({u["document_id"] for u in source_units(case)}), 3)

    def test_new_uploads_are_actually_parsed_not_filename_answers(self):
        bad = parse_case([("invoice.pdf", b"not a PDF")], storage=self.root)
        self.assertEqual(bad["documents"][0]["processing_status"], "failed")
        path = PROJECT / "data/synthetic/rfq-inbox/a-complete/inputs/items.xlsx"
        filled = parse_case([("manufacturing-rfq-template.xlsx", path.read_bytes())], storage=self.root)
        self.assertEqual(filled["documents"][0]["document_type"], "rfq")
        template = PROJECT / "data/raw/rfq-demo/manufacturing-rfq-template.xlsx"
        blank = parse_case([("arbitrary-name.xlsx", template.read_bytes())], storage=self.root)
        self.assertEqual(blank["documents"][0]["document_type"], "rfq_template")
        self.assertEqual(len(blank["documents"][0]["sheets"]), 2)

    def test_pdf_page_evidence_and_existing_document_id(self):
        path = PROJECT / "data/raw/rfq-demo/invoice-test-pack/shopify-invoice-test-pack-main/fixtures/02-multipage/invoice.pdf"
        case = parse_case([("renamed.pdf", path.read_bytes())], storage=self.root)
        doc = case["documents"][0]
        self.assertEqual(len(doc["pages"]), 2)
        evidence = {"document_id": doc["document_id"], "page": 2, "quote": doc["pages"][1]["text"][:50]}
        self.assertTrue(evidence_check(evidence, case)[0])
        self.assertFalse(evidence_check(dict(evidence, page=99), case)[0])
        existing = [json.loads(line) for line in (PROJECT / "data/processed/rfq-demo/inputs.jsonl").read_text(encoding="utf-8").splitlines()]
        original = next(d for d in existing if any(s["sha256"] == doc["source_files"][0]["sha256"] for s in d["source_files"]))
        self.assertEqual(doc["document_id"], original["document_id"])

    def test_model_evidence_gate_and_conflict_cannot_be_silently_overridden(self):
        case = self.example("c-quantity-conflict")
        draft = blank_result("rfq")
        draft.pop("origin")
        candidate = case["source_conflicts"][0]["candidates"][0]
        draft["items"] = [{k: field() for k in columns("rfq")}]
        draft["items"][0]["part_number"] = field("0007-A", candidate["evidence"][:1], "supported")
        draft["items"][0]["quantity"] = field("100", candidate["evidence"][1:], "supported")
        draft["fields"]["customer_name"] = field("Invented unsupported value", [], "supported")
        checked = verify_prediction(draft, case, "rfq")
        self.assertIsNone(checked["items"][0]["quantity"]["value"])
        self.assertEqual(checked["items"][0]["quantity"]["status"], "conflicting")
        self.assertEqual(checked["fields"]["customer_name"]["status"], "needs_review")

    def test_review_exports_use_confirmed_values_and_preserve_original(self):
        from openpyxl import load_workbook
        case = self.example("c-quantity-conflict")
        draft = blank_result("rfq")
        before = copy.deepcopy(draft)
        values = {key: "User confirmed " + key for key in draft["fields"]}
        values["quote_deadline"] = "2026-11-05"
        values["requested_delivery_date"] = "2026-12-01"
        row = {key: "" for key in columns("rfq")}
        row.update(part_number="0007-A", quantity="100", description="Spacer", unit="EA")
        result = confirm_review(draft, values, [row], "Material not supplied", case)
        exports = build_exports(case, result)
        self.assertEqual(draft, before)
        self.assertIn("quantity: 100 vs 120", exports["request-for-missing-information.txt"].decode())
        self.assertIn("material", exports["request-for-missing-information.txt"].decode())
        conflict = case["source_conflicts"][0]
        resolved = confirm_review(draft, values, [row], "Selected Excel quantity", case,
                                  {conflict["conflict_id"]: {"value": "120", "note": "Checked the source"}})
        export = build_exports(case, resolved)
        payload = json.loads(export["review.json"])
        self.assertEqual(payload["result"]["items"][0]["quantity"]["value"], "120")
        self.assertEqual(payload["result"]["fields"]["quote_deadline"]["value"], "2026-11-05")
        self.assertEqual(payload["result"]["fields"]["requested_delivery_date"]["value"], "2026-12-01")
        self.assertNotIn("conflicting information", export["request-for-missing-information.txt"].decode())
        book = load_workbook(io.BytesIO(export["items.xlsx"]), data_only=False)
        self.assertEqual(book.active["A2"].value, "0007-A")
        self.assertEqual(book.active["A2"].data_type, "s")
        save_review(case, resolved)
        self.assertTrue((Path(case["storage_dir"]) / "exports/review.json").is_file())
        with self.assertRaises(ValueError):
            build_exports(case, draft)

    def test_provider_contract_is_json_only_and_no_model_config_is_honest(self):
        case = self.example("a-complete")
        provider = OpenAICompatibleProvider({"base_url": "", "model": "", "api_key": ""})
        self.assertFalse(provider.configured())
        with self.assertRaises(ModelNotConfigured):
            provider.extract(case, "rfq")
        response = blank_result("rfq"); response.pop("origin")
        captured = []
        class Handler(BaseHTTPRequestHandler):
            def do_POST(self):
                captured.append(json.loads(self.rfile.read(int(self.headers["Content-Length"]))))
                payload = json.dumps({"choices": [{"finish_reason": "stop", "message": {"content": json.dumps(response)}}]}).encode()
                self.send_response(200); self.end_headers(); self.wfile.write(payload)
            def log_message(self, *args):
                pass
        server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True); thread.start()
        try:
            model = OpenAICompatibleProvider({"base_url": "http://127.0.0.1:" + str(server.server_port), "model": "explicit-test-double", "api_key": "", "timeout": 10})
            result = model.extract(case, "rfq")
            self.assertEqual(result["prediction"], response)
            self.assertNotIn("tools", captured[0])
            self.assertIn("untrusted data", captured[0]["messages"][0]["content"])
        finally:
            server.shutdown(); server.server_close(); thread.join(timeout=5)


if __name__ == "__main__":
    unittest.main(verbosity=2)
