#!/usr/bin/env python3
"""Validate RFQ preprocessing against the actual raw files, not just manifest claims."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import Counter
from email import policy
from email.parser import BytesParser
from email.utils import getaddresses
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from prepare_rfq_data import (PROJECT, TOOL, INPUT_KEYS, Outputs, discover, evidence_for,
                              evidence_text, json_text, json_value, local_schema, now,
                              quality_reasons, read_json, relative, render_report, sha256,
                              snapshot, stable_id)


def load_jsonl(path, errors):
    records = []
    try:
        for number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
            try:
                value = json.loads(line)
                if not isinstance(value, dict):
                    raise ValueError("Expected one JSON object")
                records.append(value)
            except Exception as exc:
                errors.append(path.name + " line " + str(number) + ": " + str(exc))
    except Exception as exc:
        errors.append("Cannot read " + str(path) + ": " + str(exc))
    return records


def validate(data_dir):
    root = Path(data_dir).resolve()
    manifest = read_json(root / "manifest.json")
    if manifest.get("tool") != TOOL:
        raise ValueError("Unrecognized processed manifest tool")
    raw = Path(manifest["raw_dir"]).resolve()
    errors, warnings = [], []
    checks = Counter()
    def check(condition, message, group):
        checks[group] += 1
        if not condition:
            errors.append(message)
    def safe_raw(name):
        path = (raw / name).resolve()
        path.relative_to(raw)
        return path
    def safe_out(name):
        path = (root / name).resolve()
        path.relative_to(root)
        return path

    inputs = load_jsonl(root / "inputs.jsonl", errors)
    gold = load_jsonl(root / "gold.jsonl", errors)
    ids = [d.get("document_id") for d in inputs]
    check(all(isinstance(i, str) and i.startswith("rfqdoc_") for i in ids), "Invalid document_id", "identity")
    check(len(ids) == len(set(ids)), "Duplicate input document_id", "identity")
    gold_ids = [g.get("document_id") for g in gold]
    check(len(gold_ids) == len(set(gold_ids)), "Duplicate gold document_id", "identity")
    by_id = {d.get("document_id"): d for d in inputs}
    primary_counts, variant_counts = Counter(), Counter()
    invoice_parents = {}
    forbidden = {"answer", "answers", "expected", "expected_json", "gold", "extraction_yaml",
                 "original_json_text", "original_yaml", "supplier", "invoiceNumber", "unitPrice",
                 "subtotal", "lineTotal", "currency", "total", "lines"}

    def no_label_keys(value, location):
        if isinstance(value, dict):
            for key, item in value.items():
                check(key not in forbidden, location + ": structured answer field " + key, "label_separation")
                no_label_keys(item, location + "/" + key)
        elif isinstance(value, list):
            for item in value:
                no_label_keys(item, location)

    for doc in inputs:
        ident = doc.get("document_id", "<missing>")
        try:
            check(set(doc) <= INPUT_KEYS, ident + ": unexpected input keys", "label_separation")
            no_label_keys(doc, ident)
            sources = doc["source_files"]
            primary = [s for s in sources if s["role"] == "primary"]
            check(len(primary) == 1 or doc["processing_status"] == "failed",
                  ident + ": expected exactly one primary file", "source_files")
            for source in sources:
                path = safe_raw(source["path"])
                check(path.is_file(), ident + ": missing input " + source["path"], "source_files")
                check(path.name not in ("expected.json", "gold.jsonl", "README.md", "CITATION.cff")
                      and path.suffix.lower() not in (".json", ".yaml", ".yml"),
                      ident + ": label/metadata file used as input", "label_separation")
                if path.is_file():
                    check(sha256(path) == source["sha256"], ident + ": input hash mismatch", "source_files")
                if source["role"] == "primary":
                    primary_counts[source["path"]] += 1
                elif source["role"] == "image_variant":
                    variant_counts[source["path"]] += 1
                    check(path.suffix.lower() == ".png", ident + ": non-PNG image variant", "pairing")
            if not primary:
                continue
            path = safe_raw(primary[0]["path"])
            if doc["document_type"] == "invoice":
                parent = relative(path.parent, raw)
                check(parent not in invoice_parents, ident + ": fixture counted twice", "pairing")
                invoice_parents[parent] = ident
                for variant in (s for s in sources if s["role"] == "image_variant"):
                    check(safe_raw(variant["path"]).parent == path.parent,
                          ident + ": variant from another fixture", "pairing")
                parts = path.parent.relative_to(raw).parts
                key = "/".join(parts[parts.index("fixtures") + 1:])
                check(ident == stable_id("salorworks:fixture:" + key), ident + ": unstable fixture ID", "identity")
            elif doc["document_type"] != "email_thread":
                check(ident == stable_id(doc["document_type"] + ":" + primary[0]["path"]),
                      ident + ": unstable document ID", "identity")
            if path.suffix.lower() == ".pdf":
                try:
                    reader = PdfReader(str(path))
                    count = len(reader.pages)
                except Exception as exc:
                    check(doc["processing_status"] == "failed", ident + ": unreadable PDF not marked failed", "pdf")
                    warnings.append(ident + ": unreadable PDF, recorded as failed: " + str(exc))
                    continue
                check(doc["pdf_page_count"] == count and len(doc["pages"]) == count,
                      ident + ": PDF page count differs from output", "pdf")
                check([p["page"] for p in doc["pages"]] == list(range(1, count + 1)),
                      ident + ": page numbering is not complete/1-based", "pdf")
                for page in doc["pages"]:
                    n = page["page"]
                    source_page = reader.pages[n - 1]
                    try:
                        native = source_page.extract_text() or ""
                        check(page["native_text"] == native, ident + ": native text not source-derived on page " + str(n), "pdf")
                    except Exception:
                        check(page["status"] != "ok", ident + ": native extraction exception hidden", "pdf")
                    check(page["source_location"] == {"path": primary[0]["path"], "page": n},
                          ident + ": missing/wrong PDF provenance", "evidence")
                    check(page["status"] in ("ok", "needs_ocr", "needs_review", "failed"), ident + ": invalid page status", "status")
                    if page["status"] == "ok" and page.get("page_type") not in ("photo", "blank"):
                        check(bool(page["text"].strip()) and not quality_reasons(page["text"]),
                              ident + ": empty/suspect text counted as success", "status")
                    else:
                        check(bool(page["warnings"]), ident + ": pending/failed page lacks reason", "status")
                    if page["extraction_method"] == "native_text":
                        check(page["text"] == page["native_text"], ident + ": native final text changed", "label_separation")
                    elif page["extraction_method"] == "ocr":
                        detail = page.get("ocr") or {}
                        check(detail.get("render_source") == primary[0]["path"] and detail.get("rendered_pdf_page") == n,
                              ident + ": OCR used another source/page (possibly first-page PNG)", "pdf")
                        if detail.get("rendered_image"):
                            image = safe_out(detail["rendered_image"])
                            check(image.is_file() and sha256(image) == detail["rendered_image_sha256"],
                                  ident + ": rendered page image missing or changed", "pdf")
                            check(page.get("ocr_raw_text") == detail.get("raw_output"), ident + ": raw OCR output not preserved", "pdf")
                    else:
                        check(False, ident + ": unknown extraction method", "pdf")
                states = [p["status"] for p in doc["pages"]]
                check(not (doc["processing_status"] == "ok" and any(s != "ok" for s in states)),
                      ident + ": document hides incomplete pages", "status")
            elif path.suffix.lower() == ".xlsx":
                formula_book = load_workbook(path, data_only=False)
                cache_book = load_workbook(path, data_only=True)
                try:
                    check([s["name"] for s in doc["sheets"]] == formula_book.sheetnames, ident + ": missing Excel sheets", "excel")
                    for record in doc["sheets"]:
                        sheet = formula_book[record["name"]]
                        cache_sheet = cache_book[record["name"]]
                        check(sorted(record["merged_ranges"]) == sorted(str(r) for r in sheet.merged_cells.ranges),
                              ident + ": merged cells not preserved", "excel")
                        check(record["max_row"] == sheet.max_row and record["max_column"] == sheet.max_column,
                              ident + ": sheet bounds changed", "excel")
                        cells = {c["coordinate"]: c for c in record["cells"]}
                        check(len(cells) == len(record["cells"]) == sheet.max_row * sheet.max_column,
                              ident + ": missing/duplicate cells", "excel")
                        for row in sheet:
                            for cell in row:
                                out = cells[cell.coordinate]
                                value = cell.value
                                formula = (value if isinstance(value, str) else getattr(value, "text", str(value))) if cell.data_type == "f" else None
                                check(out["raw_value"] == json_value(value) and out["data_type"] == cell.data_type
                                      and out["number_format"] == cell.number_format and out["formula"] == formula,
                                      ident + ": cell differs from original " + record["name"] + "!" + cell.coordinate, "excel")
                                expected_cache = json_value(cache_sheet[cell.coordinate].value) if formula is not None else None
                                check(out["cached_value"] == expected_cache,
                                      ident + ": cache changed/null replaced at " + cell.coordinate, "excel")
                                check(out["source_location"] == {"path": primary[0]["path"], "sheet": sheet.title, "cell": cell.coordinate},
                                      ident + ": cell provenance missing", "evidence")
                    check(doc["metadata"].get("template_only") is True and doc["metadata"].get("customer_data") is False,
                          ident + ": template misrepresented as customer data", "excel")
                finally:
                    formula_book.close()
                    cache_book.close()
            elif path.suffix.lower() == ".parquet":
                import pyarrow.parquet as pq
                index = doc["metadata"]["row_index"]
                original = pq.read_table(path, columns=["input"]).slice(index, 1).to_pylist()[0]["input"]
                check(doc["email_text"] == original and doc.get("task_scope") == "full_thread",
                      ident + ": email thread changed/scope lost", "parquet")
            check(read_json(safe_out(doc["text_artifact"])) == evidence_for(doc),
                  ident + ": structured text evidence differs from inputs", "evidence")
            check(safe_out(doc["text_preview_path"]).read_text(encoding="utf-8") == evidence_text(doc),
                  ident + ": plain text evidence differs from inputs", "evidence")
        except Exception as exc:
            errors.append(ident + ": validation exception: " + str(exc))

    for path, count in primary_counts.items():
        if not path.endswith(".parquet"):
            check(count == 1, "Primary file counted as multiple documents: " + path, "pairing")
    check(not (set(primary_counts) & set(variant_counts)), "PNG variant counted as a primary document", "pairing")
    fixture_dirs, pdfs, images, excels, parquets = discover(raw)
    check(len(invoice_parents) == len(fixture_dirs), "Discovered invoice fixtures missing from inputs", "pairing")
    expected_inputs = {relative(p, raw) for p in pdfs + images + excels}
    check(expected_inputs <= set(primary_counts), "Standalone source files missing from inputs", "source_files")
    complete_pairs = sum(all((d / n).is_file() for n in ("invoice.pdf", "invoice.png", "expected.json")) for d in fixture_dirs)
    check(manifest["fixture_complete_pairs"] == complete_pairs, "Incorrect complete pairing count", "pairing")

    for label in gold:
        ident = label.get("document_id")
        try:
            check(ident in by_id, "Gold record has no linked input: " + str(ident), "gold")
            source = safe_raw(label["source_file"]["path"])
            if label["answer_source"] == "publisher_expected_json":
                actual = read_json(source)
                check(label["answer"] == actual and label["original_json_text"] == source.read_text(encoding="utf-8-sig"),
                      ident + ": publisher answer structure/content changed", "gold")
                primary = next(s for s in by_id[ident]["source_files"] if s["role"] == "primary")
                check(source.parent == safe_raw(primary["path"]).parent, ident + ": wrong fixture label association", "gold")
                check(sha256(source) == label["source_file"]["sha256"], ident + ": label source hash mismatch", "gold")
                schema_file = label.get("schema_file")
                if schema_file:
                    schema = safe_raw(schema_file["path"])
                    check(sha256(schema) == schema_file["sha256"], ident + ": schema hash mismatch", "gold")
                    found = list(local_schema(schema).iter_errors(actual))
                    check(label["validation"]["status"] == ("failed" if found else "passed"),
                          ident + ": schema validation was not accurately recorded", "gold")
                    check(safe_out(label["schema_artifact"]).read_bytes() == schema.read_bytes(),
                          ident + ": schema copy differs from publisher", "gold")
                if label["validation"]["status"] != "passed":
                    warnings.append(ident + ": publisher label has not passed schema validation")
                raw_label = label["original_json_text"].strip()
                for artifact in (by_id[ident]["text_artifact"], by_id[ident]["text_preview_path"]):
                    check(not raw_label or raw_label not in safe_out(artifact).read_text(encoding="utf-8"),
                          ident + ": full label copied into extracted text", "label_separation")
            elif label["answer_source"] == "extraction_yaml":
                import yaml
                check(label.get("label_trust") == "model_generated_unverified" and label.get("eligible_for_reliable_evaluation") is False,
                      ident + ": unverified email label included in reliable evaluation", "parquet")
                # Reload the original Parquet label, not just the processed YAML copy.
                import pyarrow.parquet as pq
                original = pq.read_table(source, columns=["extraction_yaml"]).slice(label["source_file"]["row_index"], 1).to_pylist()[0]["extraction_yaml"]
                check(label["original_yaml"] == original, ident + ": original YAML changed", "parquet")
                if label["validation"]["status"] != "failed":
                    check(label["answer"] == json.loads(json.dumps(yaml.safe_load(original), default=json_value)),
                          ident + ": safe YAML result differs", "parquet")
        except Exception as exc:
            # A corrupt label is a recorded processing error, not a silently missing label.
            if label.get("validation", {}).get("status") == "failed" and label.get("answer") is None:
                warnings.append(str(ident) + ": failed label remains recorded: " + str(exc))
            else:
                errors.append(str(ident) + ": gold validation exception: " + str(exc))
    expected_gold = {invoice_parents.get(relative(d, raw)) for d in fixture_dirs if (d / "expected.json").is_file()}
    check(expected_gold <= set(gold_ids), "Existing publisher labels missing from gold.jsonl", "gold")

    for demo in manifest["demo_emails"]:
        if demo["status"] != "created":
            warnings.append("Demo email skipped: " + demo["path"] + ": " + demo.get("reason", "unknown"))
            continue
        try:
            msg = BytesParser(policy=policy.default).parsebytes(safe_out(demo["path"]).read_bytes())
            check(not msg.defects, demo["path"] + ": malformed MIME", "eml")
            check("SYNTHETIC DEMO" in str(msg["Subject"]), demo["path"] + ": missing demo subject", "eml")
            addresses = getaddresses(msg.get_all("From", []) + msg.get_all("To", []))
            check(bool(addresses) and all(address.lower().endswith("@example.com") for _, address in addresses),
                  demo["path"] + ": non-fictional address", "eml")
            attachments = list(msg.iter_attachments())
            check(len(attachments) == 1, demo["path"] + ": wrong attachment count", "eml")
            for part in attachments:
                payload = part.get_payload(decode=True)
                source = safe_raw(demo["attachment_source_path"])
                check(part.get_content_type() == "application/pdf" and part.get_filename() == source.name,
                      demo["path"] + ": wrong attachment type/name", "eml")
                check(payload == source.read_bytes() and hashlib.sha256(payload).hexdigest() == demo["attachment_sha256"],
                      demo["path"] + ": attachment bytes differ from original", "eml")
            body = msg.get_body(preferencelist=("plain",)).get_content()
            expected_body = ("Please review the attached PDF. This is historical RFQ demonstration material, not a current procurement request."
                             if "historical" in demo["path"] else "Please review the attached invoice PDF. This is a synthetic document-processing demonstration.")
            expected_body += "\nThis message is for local upload testing only.\n"
            check(body.replace("\r\n", "\n") == expected_body, demo["path"] + ": unexpected business content in body", "label_separation")
            check(demo["document_id"] in by_id, demo["path"] + ": invalid document association", "eml")
        except Exception as exc:
            errors.append(demo["path"] + ": " + str(exc))

    check(snapshot(raw) == manifest["raw_inventory"] and manifest["raw_unchanged"],
          "Raw directory changed since processing", "raw_read_only")
    manifest_docs = {d["document_id"]: d for d in manifest["documents"]}
    check(set(manifest_docs) == set(ids), "Manifest document IDs differ from inputs", "manifest")
    for doc in inputs:
        if doc["document_id"] in manifest_docs:
            for key, value in manifest_docs[doc["document_id"]].items():
                check(doc.get(key) == value, doc["document_id"] + ": manifest differs from input: " + key, "manifest")
    pending = [{"document_id": d["document_id"], "status": d["processing_status"],
                "pages": [{"page": p["page"], "status": p["status"], "source": p["source_location"]}
                          for p in d["pages"] if p["status"] != "ok"]}
               for d in inputs if d["processing_status"] != "ok"]
    if pending:
        warnings.append("Data remains incomplete: " + str(len(pending)) + " documents require OCR/review or failed; see pending_processing")
    validation = {"status": "failed" if errors else "passed", "validated_at_utc": now(),
                  "checks": dict(checks), "errors": errors, "warnings": warnings,
                  "pending_processing": pending, "processing_status_counts": dict(Counter(d["processing_status"] for d in inputs)),
                  "note": "Integrity validation passing does not mean OCR completion or business-field extraction accuracy."}
    report = render_report(manifest, inputs, gold, validation)
    # Every unresolved page/document and recorded file error must be visible to the reader.
    for item in pending:
        check(item["document_id"] in report, "Pending document absent from report", "report")
        for page in item["pages"]:
            check(page["source"]["path"] in report and ("页 " + str(page["page"])) in report,
                  "Pending page absent from report", "report")
    for issue in manifest["issues"]:
        check(issue["code"] in report, "Recorded processing issue absent from report", "report")
    validation["status"] = "failed" if errors else "passed"
    validation["checks"] = dict(checks)
    outputs = Outputs(root)
    outputs.write("validation.json", json_text(validation))
    outputs.write("processing_report.md", render_report(manifest, inputs, gold, validation))
    # Do not finish(): the validator must preserve all preparation-owned files/index entries.
    print(json.dumps({"validation": validation["status"], "checks": sum(checks.values()),
                      "errors": errors, "warnings": warnings, "pending_documents": len(pending)}, ensure_ascii=False))
    return 1 if errors else 0


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-dir", default=str(PROJECT / "data/processed/rfq-demo"))
    return validate(parser.parse_args().data_dir)


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as error:
        print("Validation aborted: " + str(error), file=sys.stderr)
        sys.exit(1)
